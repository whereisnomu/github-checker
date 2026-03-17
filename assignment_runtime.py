from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path


SUPPORTED_ASSIGNMENT_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


@dataclass(slots=True)
class AssignmentDocument:
    filename: str
    text: str
    saved_at_utc: str


class AssignmentStore:
    def __init__(self, root: str = ".cache/chat_specs") -> None:
        self.root = Path(root)
        self.root.mkdir(parents=True, exist_ok=True)

    def _path(self, chat_id: int) -> Path:
        return self.root / f"{chat_id}.json"

    def save(self, chat_id: int, filename: str, text: str) -> AssignmentDocument:
        document = AssignmentDocument(
            filename=filename,
            text=normalize_assignment_text(text),
            saved_at_utc=datetime.now(UTC).isoformat(),
        )
        self._path(chat_id).write_text(json.dumps(document.__dict__, ensure_ascii=False, indent=2), encoding="utf-8")
        return document

    def load(self, chat_id: int) -> AssignmentDocument | None:
        path = self._path(chat_id)
        if not path.exists():
            return None
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return None
        filename = str(data.get("filename") or "").strip()
        text = normalize_assignment_text(str(data.get("text") or ""))
        saved_at = str(data.get("saved_at_utc") or "").strip()
        if not filename or not text:
            return None
        return AssignmentDocument(filename=filename, text=text, saved_at_utc=saved_at)

    def clear(self, chat_id: int) -> bool:
        path = self._path(chat_id)
        if not path.exists():
            return False
        path.unlink()
        return True


def parse_assignment_document(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_ASSIGNMENT_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_ASSIGNMENT_EXTENSIONS))
        raise ValueError(f"Поддерживаются только файлы {supported}. Формат {extension or 'без расширения'} пока не поддерживается.")
    if extension == ".pdf":
        return extract_pdf_text(content)
    if extension == ".docx":
        return extract_docx_text(content)
    return normalize_assignment_text(content.decode("utf-8", errors="ignore"))


def extract_pdf_text(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("Для чтения PDF нужно установить зависимость pypdf.") from exc
    reader = PdfReader(io.BytesIO(content))
    chunks: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            chunks.append(page_text)
    text = normalize_assignment_text("\n".join(chunks))
    if not text:
        raise ValueError("Не удалось извлечь текст из PDF. Возможно, документ состоит из сканов без OCR.")
    return text


def extract_docx_text(content: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Для чтения DOCX нужно установить зависимость python-docx.") from exc
    document = Document(io.BytesIO(content))
    chunks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    text = normalize_assignment_text("\n".join(chunks))
    if not text:
        raise ValueError("Не удалось извлечь текст из DOCX.")
    return text


def normalize_assignment_text(text: str) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", " ")
    normalized = re.sub(r"[ \t]+", " ", normalized)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    normalized = normalized.strip()
    return normalized[:70000]
